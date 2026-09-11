import os
import re
import sys
import csv
import json
import math
import sqlite3
import shutil
import hashlib
import zipfile
import tempfile
import html
import threading
from pathlib import Path
from datetime import datetime, date
import tkinter as tk
from tkinter import ttk, filedialog, messagebox, simpledialog, colorchooser

try:
    from PIL import Image, ImageTk, ImageDraw, ImageFont
except Exception:
    Image = None
    ImageTk = None
    ImageDraw = None
    ImageFont = None

try:
    from openpyxl import load_workbook
except Exception:
    load_workbook = None

try:
    from docx import Document
except Exception:
    Document = None

from face_engine import FaceEngine

APP_NAME = 'Фотокаталог — Архив Президента Кыргызской Республики'
DEFAULT_PASSWORD = '12345'
IMAGE_EXTS = {'.jpg', '.jpeg', '.png', '.bmp', '.tif', '.tiff', '.webp'}


def resource_path(*parts):
    """Return a bundled-resource path both in source mode and PyInstaller EXE."""
    if getattr(sys, 'frozen', False) and hasattr(sys, '_MEIPASS'):
        base = Path(sys._MEIPASS)
    else:
        base = Path(__file__).resolve().parent
    return base.joinpath(*parts)


def app_data_dir():
    if getattr(sys, 'frozen', False):
        base = os.environ.get('LOCALAPPDATA') or os.path.expanduser('~')
        root = Path(base) / 'PhotoArchiveCatalog_Katalog5'
    else:
        root = Path(__file__).resolve().parent / 'data_katalog5'
    root.mkdir(parents=True, exist_ok=True)
    (root / 'photos').mkdir(parents=True, exist_ok=True)
    (root / 'thumbs').mkdir(parents=True, exist_ok=True)
    return root


DATA_DIR = app_data_dir()
DB_PATH = DATA_DIR / 'catalog.db'
PHOTOS_DIR = DATA_DIR / 'photos'
THUMBS_DIR = DATA_DIR / 'thumbs'


def safe_text(v):
    if v is None:
        return ''
    if isinstance(v, (datetime, date)):
        return v.strftime('%d.%m.%Y')
    return str(v).strip()


def normalize_key(s):
    return re.sub(r'[^0-9A-Za-zА-Яа-яЁё]+', '', safe_text(s)).casefold()


def whole_word_match(haystack, needle):
    haystack = safe_text(haystack)
    needle = safe_text(needle)
    if not needle:
        return True
    # Unicode-aware whole token/phrase boundaries. Hyphens in archive numbers remain searchable.
    pattern = r'(?<!\w)' + re.escape(needle) + r'(?!\w)'
    return re.search(pattern, haystack, flags=re.IGNORECASE | re.UNICODE) is not None


def copy_into_archive(src_path):
    src = Path(src_path)
    if not src.exists():
        raise FileNotFoundError(str(src))
    digest = hashlib.sha1((str(src.resolve()) + str(src.stat().st_mtime_ns)).encode('utf-8', 'ignore')).hexdigest()[:12]
    suffix = src.suffix.lower() or '.jpg'
    dest = PHOTOS_DIR / f"{datetime.now().strftime('%Y%m%d_%H%M%S_%f')}_{digest}{suffix}"
    shutil.copy2(str(src), str(dest))
    return str(dest)


def make_thumb(path, key, size=(86, 64)):
    if Image is None or not path or not os.path.exists(path):
        return None
    target = THUMBS_DIR / f'{key}_{size[0]}x{size[1]}.jpg'
    try:
        if not target.exists() or target.stat().st_mtime < Path(path).stat().st_mtime:
            img = Image.open(path)
            img.thumbnail(size)
            bg = Image.new('RGB', size, 'white')
            if img.mode != 'RGB':
                img = img.convert('RGB')
            bg.paste(img, ((size[0]-img.width)//2, (size[1]-img.height)//2))
            bg.save(str(target), 'JPEG', quality=88)
        return str(target)
    except Exception:
        return None


class CatalogDB:
    FIELDS = ('archive_no', 'description', 'shot_date', 'location', 'author', 'source', 'file_path')

    def __init__(self, path):
        self.path = str(path)
        self.conn = sqlite3.connect(self.path)
        self.conn.row_factory = sqlite3.Row
        self.conn.execute('PRAGMA foreign_keys=ON')
        self.create_schema()

    def create_schema(self):
        self.conn.executescript('''
        CREATE TABLE IF NOT EXISTS photos (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            archive_no TEXT NOT NULL DEFAULT '',
            description TEXT NOT NULL DEFAULT '',
            shot_date TEXT NOT NULL DEFAULT '',
            location TEXT NOT NULL DEFAULT '',
            author TEXT NOT NULL DEFAULT '',
            source TEXT NOT NULL DEFAULT '',
            file_path TEXT NOT NULL DEFAULT '',
            created_at TEXT NOT NULL,
            updated_at TEXT NOT NULL
        );
        CREATE INDEX IF NOT EXISTS idx_photos_archive_no ON photos(archive_no);
        CREATE INDEX IF NOT EXISTS idx_photos_shot_date ON photos(shot_date);
        CREATE INDEX IF NOT EXISTS idx_photos_author ON photos(author);
        ''')
        # v3 face index: several faces can belong to one photo.
        cols = [r['name'] for r in self.conn.execute("PRAGMA table_info(face_index)").fetchall()]
        if cols and 'id' not in cols:
            self.conn.execute('DROP TABLE IF EXISTS face_index')
        self.conn.executescript('''
        CREATE TABLE IF NOT EXISTS face_index (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            photo_id INTEGER NOT NULL,
            face_no INTEGER NOT NULL DEFAULT 0,
            embedding TEXT NOT NULL,
            indexed_at TEXT NOT NULL,
            FOREIGN KEY(photo_id) REFERENCES photos(id) ON DELETE CASCADE
        );
        CREATE INDEX IF NOT EXISTS idx_face_photo ON face_index(photo_id);
        ''')
        self.conn.commit()

    def add(self, rec, commit=True):
        now = datetime.now().isoformat(timespec='seconds')
        cur = self.conn.execute('''INSERT INTO photos
        (archive_no,description,shot_date,location,author,source,file_path,created_at,updated_at)
        VALUES (?,?,?,?,?,?,?,?,?)''',
        (rec.get('archive_no',''), rec.get('description',''), rec.get('shot_date',''), rec.get('location',''),
         rec.get('author',''), rec.get('source',''), rec.get('file_path',''), now, now))
        if commit:
            self.conn.commit()
        return cur.lastrowid

    def update(self, photo_id, rec, commit=True, reset_faces=True):
        now = datetime.now().isoformat(timespec='seconds')
        self.conn.execute('''UPDATE photos SET archive_no=?,description=?,shot_date=?,location=?,author=?,source=?,file_path=?,updated_at=? WHERE id=?''',
                          (rec.get('archive_no',''),rec.get('description',''),rec.get('shot_date',''),rec.get('location',''),
                           rec.get('author',''),rec.get('source',''),rec.get('file_path',''),now,photo_id))
        if reset_faces:
            self.conn.execute('DELETE FROM face_index WHERE photo_id=?',(photo_id,))
        if commit:
            self.conn.commit()

    def archive_index(self):
        rows = self.conn.execute('SELECT * FROM photos ORDER BY id').fetchall()
        out = {}
        for r in rows:
            out.setdefault(safe_text(r['archive_no']).casefold(), []).append(r)
        return out

    def bulk_upsert(self, records, progress_cb=None, batch_size=1000):
        """Fast metadata import: one lookup cache + batched commits."""
        cache = self.archive_index()
        added = updated = skipped = 0
        total = len(records)
        self.conn.execute('BEGIN')
        try:
            for i, rec in enumerate(records, 1):
                archive_no = safe_text(rec.get('archive_no'))
                if not archive_no:
                    skipped += 1
                else:
                    key = archive_no.casefold()
                    existing = cache.get(key, [])
                    if existing:
                        old = existing[0]
                        merged = {k: (rec.get(k) if safe_text(rec.get(k)) else old[k]) for k in self.FIELDS}
                        self.update(old['id'], merged, commit=False, reset_faces=False)
                        updated += 1
                        # keep cache current for subsequent duplicate rows in the same import
                        fresh = self.get(old['id'])
                        cache[key][0] = fresh
                    else:
                        new_id = self.add(rec, commit=False)
                        fresh = self.get(new_id)
                        cache.setdefault(key, []).append(fresh)
                        added += 1
                if i % batch_size == 0:
                    self.conn.commit()
                    self.conn.execute('BEGIN')
                if progress_cb and (i == total or i % 50 == 0):
                    progress_cb(i, total, added, updated, skipped)
            self.conn.commit()
        except Exception:
            self.conn.rollback()
            raise
        return added, updated, skipped

    def delete(self, photo_id):
        row = self.get(photo_id)
        self.conn.execute('DELETE FROM face_index WHERE photo_id=?',(photo_id,))
        self.conn.execute('DELETE FROM photos WHERE id=?',(photo_id,))
        self.conn.commit(); return row

    def get(self, photo_id):
        return self.conn.execute('SELECT * FROM photos WHERE id=?',(photo_id,)).fetchone()

    def all(self):
        return self.conn.execute('SELECT * FROM photos ORDER BY id DESC').fetchall()

    def by_archive_no(self, archive_no):
        return self.conn.execute('SELECT * FROM photos WHERE lower(archive_no)=lower(?) ORDER BY id',(archive_no,)).fetchall()

    def search(self, text='', field='Все поля', whole_word=True):
        text = safe_text(text)
        rows = self.all()
        if not text:
            return rows
        fmap = {
            'Архивный номер': ['archive_no'], 'Описание': ['description'], 'Дата съёмки': ['shot_date'],
            'Место съёмки': ['location'], 'Автор': ['author'], 'Источник': ['source'],
            'Все поля': ['archive_no','description','shot_date','location','author','source']
        }
        fields = fmap.get(field, fmap['Все поля'])
        if whole_word:
            terms = [t for t in re.split(r'\s+', text) if t]
            result = []
            for r in rows:
                combined = ' '.join(safe_text(r[f]) for f in fields)
                if all(whole_word_match(combined, t) for t in terms):
                    result.append(r)
            return result
        low = text.casefold()
        return [r for r in rows if any(low in safe_text(r[f]).casefold() for f in fields)]

    def replace_faces(self, photo_id, embeddings):
        now = datetime.now().isoformat(timespec='seconds')
        self.conn.execute('DELETE FROM face_index WHERE photo_id=?',(photo_id,))
        for i, emb in enumerate(embeddings):
            payload = json.dumps([float(x) for x in emb])
            self.conn.execute('INSERT INTO face_index(photo_id,face_no,embedding,indexed_at) VALUES (?,?,?,?)',
                              (photo_id,i,payload,now))
        self.conn.commit()

    def embeddings(self):
        return self.conn.execute('''SELECT f.id face_id,f.photo_id,f.face_no,f.embedding,p.*
                                    FROM face_index f JOIN photos p ON p.id=f.photo_id''').fetchall()

    def close(self): self.conn.close()


class PasswordDialog(simpledialog.Dialog):
    def body(self, master):
        ttk.Label(master, text='Введите пароль для перехода в защищённый раздел:').grid(row=0,column=0,padx=8,pady=(8,5))
        self.var = tk.StringVar()
        ent = ttk.Entry(master, textvariable=self.var, show='*', width=28)
        ent.grid(row=1,column=0,padx=8,pady=(0,8)); ent.focus_set(); return ent
    def apply(self): self.result = self.var.get()


class PhotoPopup(tk.Toplevel):
    def __init__(self, master, path, title='Фотография', max_size=(1100,760)):
        super().__init__(master); self.title(title); self.transient(master)
        self.ref = None
        if not path or not os.path.exists(path) or Image is None:
            ttk.Label(self,text='Файл фотографии не найден.',padding=30).pack(); return
        try:
            img=Image.open(path); img.thumbnail(max_size)
            self.ref=ImageTk.PhotoImage(img)
            ttk.Label(self,image=self.ref).pack(padx=8,pady=8)
        except Exception as exc:
            ttk.Label(self,text=f'Не удалось открыть фото:\n{exc}',padding=30).pack()


class Splash(tk.Toplevel):
    def __init__(self, master):
        super().__init__(master)
        self.overrideredirect(True)
        self.configure(bg='white')
        self._image_refs = []
        w, h = 820, 560
        sw, sh = self.winfo_screenwidth(), self.winfo_screenheight()
        self.geometry(f'{w}x{h}+{(sw-w)//2}+{(sh-h)//2}')

        canvas = tk.Canvas(self, width=w, height=h, bg='#f4f8fb',
                           highlightthickness=1, highlightbackground='#174e70')
        canvas.pack(fill='both', expand=True)

        # Clean top area: real flag and state emblem are bundled into the EXE.
        canvas.create_rectangle(0, 0, w, 210, fill='white', outline='')
        flag_path = resource_path('assets', 'flag_kg.png')
        emblem_path = resource_path('assets', 'emblem_kg.png')

        if Image is not None and flag_path.exists():
            try:
                flag = Image.open(flag_path).convert('RGBA')
                flag.thumbnail((330, 175))
                flag_ref = ImageTk.PhotoImage(flag)
                self._image_refs.append(flag_ref)
                canvas.create_image(235, 105, image=flag_ref)
            except Exception:
                pass

        if Image is not None and emblem_path.exists():
            try:
                emblem = Image.open(emblem_path).convert('RGBA')
                emblem.thumbnail((170, 170))
                emblem_ref = ImageTk.PhotoImage(emblem)
                self._image_refs.append(emblem_ref)
                canvas.create_image(585, 105, image=emblem_ref)
            except Exception:
                pass

        # Fallback is intentionally text-only: no fake coat-of-arms shapes or emoji.
        if not self._image_refs:
            canvas.create_rectangle(115, 45, 365, 165, fill='#ff0000', outline='')
            canvas.create_text(240, 105, text='КЫРГЫЗ РЕСПУБЛИКАСЫ',
                               fill='#ffff00', font=('Arial', 15, 'bold'))

        canvas.create_line(70, 210, w-70, 210, fill='#c9d7e2', width=1)
        canvas.create_text(w//2, 285,
                           text='АРХИВ ПРЕЗИДЕНТА\nКЫРГЫЗСКОЙ РЕСПУБЛИКИ',
                           fill='#123a58', font=('Arial', 28, 'bold'), justify='center')
        canvas.create_text(w//2, 380, text='ФОТОКАТАЛОГ',
                           fill='#16649a', font=('Arial', 19, 'bold'))

        self.pb = ttk.Progressbar(self, mode='indeterminate', length=500)
        self.pb.place(x=160, y=455)
        self.pb.start(12)
        canvas.create_text(w//2, 500, text='Загрузка программы…',
                           fill='#374957', font=('Arial', 12))
        canvas.create_text(w//2, 535, text='Сохранение • учёт • поиск архивных фотодокументов',
                           fill='#6a7b87', font=('Arial', 10))


class ProgressDialog(tk.Toplevel):
    def __init__(self, master, title='Выполнение операции'):
        super().__init__(master)
        self.title(title)
        self.transient(master)
        self.resizable(False, False)
        self.protocol('WM_DELETE_WINDOW', lambda: None)
        self.geometry('560x190')
        self.grab_set()
        frm = ttk.Frame(self, padding=16)
        frm.pack(fill='both', expand=True)
        self.title_var = tk.StringVar(value=title)
        self.percent_var = tk.StringVar(value='0%')
        self.detail_var = tk.StringVar(value='Подготовка…')
        self.stats_var = tk.StringVar(value='')
        ttk.Label(frm, textvariable=self.title_var, font=('Arial', 12, 'bold')).pack(anchor='w')
        row = ttk.Frame(frm); row.pack(fill='x', pady=(14, 4))
        self.pb = ttk.Progressbar(row, mode='determinate', maximum=100)
        self.pb.pack(side='left', fill='x', expand=True)
        ttk.Label(row, textvariable=self.percent_var, width=7, anchor='e', font=('Arial', 11, 'bold')).pack(side='left', padx=(8, 0))
        ttk.Label(frm, textvariable=self.detail_var, font=('Arial', 10)).pack(anchor='w', pady=(6, 2))
        ttk.Label(frm, textvariable=self.stats_var, font=('Arial', 10)).pack(anchor='w')
        self.update_idletasks()

    def update_progress(self, current, total, detail='', stats=''):
        total = max(1, int(total or 1))
        current = min(total, max(0, int(current or 0)))
        pct = int(round(current * 100 / total))
        self.pb['value'] = pct
        self.percent_var.set(f'{pct}%')
        self.detail_var.set(detail or f'Обработано: {current} из {total}')
        self.stats_var.set(stats)
        self.update_idletasks()

    def close(self):
        try:
            self.grab_release()
        except Exception:
            pass
        try:
            self.destroy()
        except Exception:
            pass


class SettingsDialog(tk.Toplevel):
    def __init__(self, app):
        super().__init__(app)
        self.app = app
        self.title('Настройки программы')
        self.transient(app)
        self.resizable(False, False)
        self.geometry('470x355')
        self.mode = tk.StringVar(value=app.settings.get('mode', 'day'))
        self.bg = tk.StringVar(value=app.settings.get('background', '#eef2f5'))
        self.button = tk.StringVar(value=app.settings.get('button', '#d6dce1'))
        self.selection = tk.StringVar(value=app.settings.get('selection', '#b7d7f0'))
        frm = ttk.Frame(self, padding=16); frm.pack(fill='both', expand=True)
        ttk.Label(frm, text='Режим оформления', font=('Arial', 11, 'bold')).grid(row=0, column=0, sticky='w')
        modes = ttk.Frame(frm); modes.grid(row=1, column=0, columnspan=2, sticky='w', pady=(6, 16))
        ttk.Radiobutton(modes, text='Дневной', variable=self.mode, value='day', command=self._preset).pack(side='left')
        ttk.Radiobutton(modes, text='Ночной', variable=self.mode, value='night', command=self._preset).pack(side='left', padx=18)
        self._color_row(frm, 2, 'Цвет фона', self.bg)
        self._color_row(frm, 3, 'Цвет кнопок', self.button)
        self._color_row(frm, 4, 'Цвет выделения', self.selection)
        ttk.Separator(frm).grid(row=5, column=0, columnspan=2, sticky='ew', pady=18)
        btns = ttk.Frame(frm); btns.grid(row=6, column=0, columnspan=2, sticky='ew')
        ttk.Button(btns, text='Вернуть стандартные', command=self._reset).pack(side='left')
        ttk.Button(btns, text='Отмена', command=self.destroy).pack(side='right')
        ttk.Button(btns, text='Сохранить настройки', command=self._save).pack(side='right', padx=6)
        frm.columnconfigure(0, weight=1)

    def _color_row(self, frm, row, label, var):
        ttk.Label(frm, text=label).grid(row=row, column=0, sticky='w', pady=6)
        ttk.Button(frm, textvariable=var, command=lambda v=var:self._choose(v), width=16).grid(row=row, column=1, sticky='e', pady=6)

    def _choose(self, var):
        c = colorchooser.askcolor(var.get(), parent=self)[1]
        if c:
            var.set(c)

    def _preset(self):
        if self.mode.get() == 'night':
            self.bg.set('#25292e'); self.button.set('#3a4047'); self.selection.set('#3f6685')
        else:
            self.bg.set('#eef2f5'); self.button.set('#d6dce1'); self.selection.set('#b7d7f0')

    def _reset(self):
        self.mode.set('day'); self._preset()

    def _save(self):
        self.app.settings = {
            'mode': self.mode.get(),
            'background': self.bg.get(),
            'button': self.button.get(),
            'selection': self.selection.get()
        }
        self.app.save_settings()
        self.app.apply_theme()
        self.destroy()


class PhotoCatalogApp(tk.Tk):
    def __init__(self):
        super().__init__(); self.withdraw(); self.title(APP_NAME); self.geometry('1420x860'); self.minsize(1080,680)
        self.db=CatalogDB(DB_PATH); self.face=FaceEngine(); self.current_screen='viewer'; self.session_unlocked=False
        self.hover_popup=None; self.hover_after=None; self.current_rows=[]
        self.protocol('WM_DELETE_WINDOW',self.on_close)
        self.style=ttk.Style(self)
        try:
            self.style.theme_use('clam')
        except Exception:
            pass
        self.settings_path = DATA_DIR / 'settings.json'
        self.settings = self.load_settings()
        self.apply_theme()
        self.container=ttk.Frame(self); self.container.pack(fill='both',expand=True)
        self.screens={}
        for name, cls in [('viewer',ViewerScreen),('editor',EditorScreen),('catalog',CatalogScreen)]:
            f=cls(self.container,self); self.screens[name]=f; f.grid(row=0,column=0,sticky='nsew')
        self.container.rowconfigure(0,weight=1); self.container.columnconfigure(0,weight=1)
        splash=Splash(self)
        self.after(1400,lambda:self._finish_splash(splash))

    def load_settings(self):
        defaults = {'mode':'day','background':'#eef2f5','button':'#d6dce1','selection':'#b7d7f0'}
        try:
            if self.settings_path.exists():
                data = json.loads(self.settings_path.read_text(encoding='utf-8'))
                defaults.update({k:v for k,v in data.items() if k in defaults and v})
        except Exception:
            pass
        return defaults

    def save_settings(self):
        try:
            self.settings_path.write_text(json.dumps(self.settings, ensure_ascii=False, indent=2), encoding='utf-8')
        except Exception:
            pass

    def apply_theme(self):
        mode = self.settings.get('mode','day')
        bg = self.settings.get('background','#eef2f5')
        btn = self.settings.get('button','#d6dce1')
        sel = self.settings.get('selection','#b7d7f0')
        fg = '#f2f2f2' if mode == 'night' else '#111111'
        field = '#343a40' if mode == 'night' else '#ffffff'
        alt = '#30363c' if mode == 'night' else '#eeeeee'
        heading = '#444b52' if mode == 'night' else '#d0d5da'
        self.theme_colors = {'bg':bg,'button':btn,'selection':sel,'fg':fg,'field':field,'alt':alt,'heading':heading}
        self.configure(bg=bg)
        self.style.configure('.', font=('Arial', 10), background=bg, foreground=fg)
        self.style.configure('TFrame', background=bg)
        self.style.configure('TLabel', background=bg, foreground=fg)
        self.style.configure('TLabelframe', background=bg, foreground=fg)
        self.style.configure('TLabelframe.Label', background=bg, foreground=fg, font=('Arial',10,'bold'))
        self.style.configure('TButton', background=btn, foreground=fg, padding=(8,5), relief='raised', borderwidth=1)
        self.style.map('TButton', background=[('active', self._shade(btn, -16)), ('pressed', self._shade(btn, -28))])
        self.style.configure('TEntry', fieldbackground=field, foreground=fg)
        self.style.configure('TCombobox', fieldbackground=field, foreground=fg)
        self.style.configure('Archive.Treeview', font=('Arial', 11), rowheight=32,
                             background=field, fieldbackground=field, foreground=fg,
                             borderwidth=1, relief='solid')
        self.style.configure('Archive.Treeview.Heading', font=('Arial', 11, 'bold'),
                             background=heading, foreground=fg, relief='raised', borderwidth=1, padding=(6, 5))
        self.style.map('Archive.Treeview', background=[('selected', sel)], foreground=[('selected', fg)])
        self.style.configure('Catalog.Treeview', font=('Arial', 11), rowheight=80,
                             background=field, fieldbackground=field, foreground=fg,
                             borderwidth=1, relief='solid')
        self.style.configure('Catalog.Treeview.Heading', font=('Arial', 11, 'bold'),
                             background=heading, foreground=fg, relief='raised', borderwidth=1, padding=(6, 5))
        self.style.map('Catalog.Treeview', background=[('selected', sel)], foreground=[('selected', fg)])
        # Add a visible border around each row. Together with the vertical separators
        # inserted in displayed values, this gives the catalog a clear grid look.
        try:
            item_layout=[('Treeitem.border', {'sticky':'nswe','children':[
                ('Treeitem.padding', {'sticky':'nswe','children':[
                    ('Treeitem.indicator', {'side':'left','sticky':''}),
                    ('Treeitem.image', {'side':'left','sticky':''}),
                    ('Treeitem.text', {'side':'left','sticky':''})
                ]})
            ]})]
            self.style.layout('Archive.Treeview.Item', item_layout)
            self.style.layout('Catalog.Treeview.Item', item_layout)
            self.style.configure('Archive.Treeview.Item', borderwidth=1, relief='solid')
            self.style.configure('Catalog.Treeview.Item', borderwidth=1, relief='solid')
        except Exception:
            pass
        # Apply colors to existing classic Text widgets.
        def walk(w):
            try:
                if isinstance(w, tk.Text):
                    w.configure(bg=field, fg=fg, insertbackground=fg, selectbackground=sel)
            except Exception:
                pass
            for c in w.winfo_children():
                walk(c)
        try: walk(self)
        except Exception: pass
        for scr in getattr(self,'screens',{}).values():
            try: scr.refresh_row_colors()
            except Exception: pass

    def _shade(self, color, amount):
        try:
            color=color.lstrip('#'); vals=[int(color[i:i+2],16) for i in (0,2,4)]
            vals=[max(0,min(255,v+amount)) for v in vals]
            return '#%02x%02x%02x'%tuple(vals)
        except Exception:
            return color

    def open_settings(self):
        SettingsDialog(self)

    def _finish_splash(self,splash):
        try: splash.destroy()
        except Exception: pass
        self.deiconify(); self.show_screen('viewer', protected=False)

    def unlock(self):
        dlg=PasswordDialog(self,title='Ввод пароля')
        if dlg.result is None: return False
        if dlg.result != DEFAULT_PASSWORD:
            messagebox.showerror(APP_NAME,'Неверный пароль.'); return False
        self.session_unlocked=True; return True

    def show_screen(self,name,protected=None):
        if protected is None: protected = name in ('editor','catalog')
        if protected and not self.session_unlocked and not self.unlock(): return
        self.current_screen=name; self.screens[name].refresh(); self.screens[name].tkraise()

    def reset_lock(self): self.session_unlocked=False; self.show_screen('viewer',False)

    def popup_photo(self,path,title='Фотография'):
        PhotoPopup(self,path,title)

    def on_close(self):
        try:self.db.close()
        except Exception:pass
        self.destroy()


class BaseScreen(ttk.Frame):
    def __init__(self,parent,app):
        super().__init__(parent); self.app=app; self.db=app.db; self.thumb_refs={}; self.rows=[]
        self.status_var=tk.StringVar(value='Готово')

    def nav(self, active):
        bar=ttk.Frame(self,padding=(10,8)); bar.pack(fill='x')
        ttk.Label(bar,text='АРХИВ ПРЕЗИДЕНТА КЫРГЫЗСКОЙ РЕСПУБЛИКИ',font=('Arial',12,'bold')).pack(side='left',padx=(0,20))
        ttk.Button(bar,text='1. Просмотр и поиск',command=lambda:self.app.show_screen('viewer',False)).pack(side='left',padx=3)
        ttk.Button(bar,text='2. Редактирование 🔒',command=lambda:self.app.show_screen('editor')).pack(side='left',padx=3)
        ttk.Button(bar,text='3. Каталог 🔒',command=lambda:self.app.show_screen('catalog')).pack(side='left',padx=3)
        ttk.Button(bar,text='Заблокировать',command=self.app.reset_lock).pack(side='right',padx=3)

    def status(self):
        ttk.Label(self,textvariable=self.status_var,relief='sunken',anchor='w',padding=(6,3)).pack(fill='x',side='bottom')

    def refresh_row_colors(self):
        tree = getattr(self, 'tree', None)
        if tree is not None:
            c = self.app.theme_colors
            tree.tag_configure('even', background=c['field'], foreground=c['fg'])
            tree.tag_configure('odd', background=c['alt'], foreground=c['fg'])

    def refresh(self): pass


class ViewerScreen(BaseScreen):
    def __init__(self,parent,app):
        super().__init__(parent,app); self.nav('viewer')
        body=ttk.Panedwindow(self,orient='horizontal'); body.pack(fill='both',expand=True,padx=10,pady=(0,8))
        left=ttk.Frame(body,padding=6); right=ttk.Frame(body,padding=8); body.add(left,weight=2); body.add(right,weight=3)
        search=ttk.LabelFrame(left,text='Поиск',padding=8); search.pack(fill='x')
        self.field=tk.StringVar(value='Все поля'); ttk.Combobox(search,textvariable=self.field,state='readonly',values=['Все поля','Архивный номер','Описание','Дата съёмки','Место съёмки','Автор','Источник']).pack(fill='x',pady=(0,6))
        self.q=tk.StringVar(); ent=ttk.Entry(search,textvariable=self.q); ent.pack(fill='x'); ent.bind('<Return>',lambda e:self.refresh())
        self.whole=tk.BooleanVar(value=True); ttk.Checkbutton(search,text='Искать только целое слово',variable=self.whole).pack(anchor='w',pady=6)
        bb=ttk.Frame(search);bb.pack(fill='x');ttk.Button(bb,text='Найти',command=self.refresh).pack(side='left');ttk.Button(bb,text='Сбросить',command=self.reset).pack(side='left',padx=5)
        ttk.Button(search,text='Поиск по лицу',command=self.search_face).pack(fill='x',pady=(8,0))
        ttk.Button(search,text='🖨 Печать результата',command=self.print_results).pack(fill='x',pady=(6,0))

        listbox=ttk.LabelFrame(left,text='Результаты',padding=4); listbox.pack(fill='both',expand=True,pady=(8,0))
        self.tree=ttk.Treeview(listbox,columns=('seq','no','date','desc'),show='headings',selectmode='browse',style='Archive.Treeview')
        for c,t,w in [('seq','№',55),('no','Архивный №',130),('date','Дата',100),('desc','Описание',360)]:self.tree.heading(c,text=t);self.tree.column(c,width=w,anchor='center' if c=='seq' else 'w')
        ys=ttk.Scrollbar(listbox,orient='vertical',command=self.tree.yview); xs=ttk.Scrollbar(listbox,orient='horizontal',command=self.tree.xview)
        self.tree.configure(yscrollcommand=ys.set,xscrollcommand=xs.set); self.tree.grid(row=0,column=0,sticky='nsew');ys.grid(row=0,column=1,sticky='ns');xs.grid(row=1,column=0,sticky='ew');listbox.rowconfigure(0,weight=1);listbox.columnconfigure(0,weight=1)
        self.tree.bind('<<TreeviewSelect>>',self.select)

        self.preview=ttk.Label(right,text='Выберите запись',anchor='center'); self.preview.pack(fill='both',expand=True)
        self.preview_ref=None; self.info=tk.Text(right,height=20,wrap='word',state='disabled',font=('Arial',12),spacing1=3,spacing3=4); self.info.pack(fill='both',expand=True,pady=(8,0))
        ttk.Button(right,text='Открыть фотографию крупно',command=self.open_current).pack(anchor='e',pady=6)
        self.current=None; self.status()

    def reset(self): self.q.set('');self.refresh()
    def refresh(self):
        self.rows=self.db.search(self.q.get(),self.field.get(),self.whole.get());
        for i in self.tree.get_children(): self.tree.delete(i)
        self.refresh_row_colors()
        for idx, r in enumerate(self.rows):
            desc=safe_text(r['description']).replace('\n',' ')
            self.tree.insert('', 'end',iid=str(r['id']),values=(idx+1,'│ '+safe_text(r['archive_no']),'│ '+safe_text(r['shot_date']),'│ '+desc[:120]),
                             tags=('even' if idx % 2 == 0 else 'odd',))
        self.status_var.set(f'Найдено записей: {len(self.rows)}')
        if self.rows:
            self.tree.selection_set(str(self.rows[0]['id'])); self.select()

    def select(self,_e=None):
        s=self.tree.selection();
        if not s:return
        r=self.db.get(int(s[0])); self.current=r
        self.show_photo(r['file_path'])
        txt=(f"Архивный номер: {r['archive_no']}\nДата съёмки: {r['shot_date']}\nМесто съёмки: {r['location']}\n"
             f"Автор: {r['author']}\nИсточник поступления: {r['source']}\n\nОписание:\n{r['description']}")
        self.info.configure(state='normal'); self.info.delete('1.0','end');self.info.insert('1.0',txt);self.info.configure(state='disabled')

    def show_photo(self,path):
        self.preview_ref=None
        if not path or not os.path.exists(path) or Image is None:self.preview.configure(image='',text='Фотография отсутствует');return
        try:
            img=Image.open(path);img.thumbnail((760,500));self.preview_ref=ImageTk.PhotoImage(img);self.preview.configure(image=self.preview_ref,text='')
        except Exception:self.preview.configure(image='',text='Не удалось открыть фотографию')
    def open_current(self):
        if self.current:self.app.popup_photo(self.current['file_path'],self.current['archive_no'])

    def print_results(self):
        if not self.rows:
            messagebox.showinfo(APP_NAME, 'Нет результатов для печати.')
            return
        one = messagebox.askyesnocancel(APP_NAME,
            'Что подготовить к печати?\n\nДа — текущую карточку\nНет — все результаты поиска\nОтмена — не печатать')
        if one is None:
            return
        rows = [self.current] if one and self.current else self.rows
        if not rows:
            return
        parts = ["<html><head><meta charset='utf-8'><title>Фотокаталог</title>",
                 "<style>body{font-family:Arial,sans-serif;font-size:12pt;color:#111} h1{font-size:18pt} .card{page-break-inside:avoid;border:1px solid #777;padding:12px;margin:10px 0} img{max-width:520px;max-height:360px} table{border-collapse:collapse;width:100%} td{border:1px solid #999;padding:5px;vertical-align:top}</style></head><body>",
                 "<h1>Архив Президента Кыргызской Республики — Фотокаталог</h1>"]
        for n, r in enumerate(rows, 1):
            parts.append("<div class='card'>")
            parts.append(f"<b>№ {n}</b><br><b>Архивный номер:</b> {html.escape(safe_text(r['archive_no']))}<br>")
            fp = safe_text(r['file_path'])
            if fp and os.path.exists(fp):
                try:
                    uri = Path(fp).resolve().as_uri()
                    parts.append(f"<p><img src='{uri}'></p>")
                except Exception:
                    pass
            parts.append("<table>")
            fields=[('Дата съёмки','shot_date'),('Место съёмки','location'),('Автор','author'),('Источник поступления','source'),('Описание','description')]
            for label,key in fields:
                parts.append(f"<tr><td style='width:190px'><b>{label}</b></td><td>{html.escape(safe_text(r[key])).replace(chr(10),'<br>')}</td></tr>")
            parts.append("</table></div>")
        parts.append("</body></html>")
        out = Path(tempfile.gettempdir()) / 'PhotoArchiveCatalog_print.html'
        out.write_text(''.join(parts), encoding='utf-8')
        try:
            os.startfile(str(out))
            messagebox.showinfo(APP_NAME, 'Предварительный просмотр открыт в браузере.\nДля печати нажмите Ctrl+P и выберите принтер.')
        except Exception as exc:
            messagebox.showerror(APP_NAME, f'Не удалось открыть предпросмотр:\n{exc}')

    def search_face(self):
        if not self.app.face.available: messagebox.showerror(APP_NAME,'Поиск по лицу недоступен:\n'+self.app.face.error);return
        path=filedialog.askopenfilename(title='Выберите фотографию лица',filetypes=[('Изображения','*.jpg *.jpeg *.png *.bmp *.tif *.tiff')]);
        if not path:return
        try:q=self.app.face.embedding(path)
        except Exception as exc:messagebox.showerror(APP_NAME,str(exc));return
        best={}
        for r in self.db.embeddings():
            try:
                score=self.app.face.cosine(q,json.loads(r['embedding'])); best[r['photo_id']]=max(best.get(r['photo_id'],-1),score)
            except Exception:pass
        ids=[pid for pid,sc in sorted(best.items(),key=lambda x:x[1],reverse=True) if sc>=0.36][:200]
        self.rows=[self.db.get(pid) for pid in ids]
        for i in self.tree.get_children():self.tree.delete(i)
        self.refresh_row_colors()
        for idx,r in enumerate(self.rows):self.tree.insert('', 'end',iid=str(r['id']),values=(idx+1,'│ '+safe_text(r['archive_no']),'│ '+safe_text(r['shot_date']),'│ '+safe_text(r['description'])[:120]),tags=('even' if idx % 2 == 0 else 'odd',))
        self.status_var.set(f'Поиск по лицу: найдено {len(self.rows)}')
        if not ids:messagebox.showinfo(APP_NAME,'Совпадений нет. В окне редактирования сначала нажмите «Индексировать лица».')


class EditorScreen(BaseScreen):
    def __init__(self,parent,app):
        super().__init__(parent,app);self.nav('editor');self.current_id=None;self.current_path='';self.preview_ref=None
        top=ttk.Frame(self,padding=(10,0,10,6));top.pack(fill='x')
        ttk.Button(top,text='Новая карточка',command=self.new).pack(side='left')
        ttk.Button(top,text='Импорт Access',command=self.import_access).pack(side='left',padx=4)
        ttk.Button(top,text='Импорт Excel/CSV',command=self.import_table).pack(side='left',padx=4)
        ttk.Button(top,text='Импорт Word DOCX',command=self.import_word).pack(side='left',padx=4)
        ttk.Button(top,text='Импорт фото по архивным №',command=self.import_photos_by_number).pack(side='left',padx=12)
        ttk.Button(top,text='Индексировать лица',command=self.index_faces).pack(side='left',padx=4)
        ttk.Button(top,text='⚙ Настройки программы',command=self.app.open_settings).pack(side='right',padx=(6,0))
        ttk.Button(top,text='Резервная копия',command=self.backup).pack(side='right')
        body=ttk.Panedwindow(self,orient='horizontal');body.pack(fill='both',expand=True,padx=10,pady=(0,8))
        lf=ttk.Frame(body,padding=5);rf=ttk.Frame(body,padding=8);body.add(lf,weight=2);body.add(rf,weight=3)
        self.q=tk.StringVar(); sr=ttk.Frame(lf);sr.pack(fill='x');ttk.Entry(sr,textvariable=self.q).pack(side='left',fill='x',expand=True);ttk.Button(sr,text='Найти',command=self.refresh).pack(side='left',padx=4)
        self.tree=ttk.Treeview(lf,columns=('seq','no','date','desc'),show='headings',selectmode='browse',style='Archive.Treeview')
        for c,t,w in [('seq','№',55),('no','Архивный №',130),('date','Дата',95),('desc','Описание',350)]:self.tree.heading(c,text=t);self.tree.column(c,width=w,anchor='center' if c=='seq' else 'w')
        ys=ttk.Scrollbar(lf,orient='vertical',command=self.tree.yview);self.tree.configure(yscrollcommand=ys.set);self.tree.pack(side='left',fill='both',expand=True,pady=(6,0));ys.pack(side='right',fill='y',pady=(6,0));self.tree.bind('<<TreeviewSelect>>',self.select)
        self.entries={}; form=rf
        row=0
        for k,label in [('archive_no','Архивный номер'),('shot_date','Дата съёмки'),('location','Место съёмки'),('author','Автор съёмки'),('source','Источник поступления')]:
            ttk.Label(form,text=label).grid(row=row,column=0,sticky='w');e=ttk.Entry(form);e.grid(row=row+1,column=0,sticky='ew',pady=(0,6));self.entries[k]=e;row+=2
        ttk.Label(form,text='Описание фотографии').grid(row=row,column=0,sticky='w');row+=1;self.desc=tk.Text(form,height=6,wrap='word');self.desc.grid(row=row,column=0,sticky='nsew',pady=(0,6));row+=1
        pf=ttk.LabelFrame(form,text='Фотография',padding=6);pf.grid(row=row,column=0,sticky='nsew');self.preview=ttk.Label(pf,text='Фото не выбрано',anchor='center');self.preview.pack(fill='both',expand=True)
        pb=ttk.Frame(pf);pb.pack(fill='x',pady=5);ttk.Button(pb,text='Выбрать/заменить фото',command=self.choose_photo).pack(side='left');ttk.Button(pb,text='Удалить фото',command=self.remove_photo).pack(side='left',padx=5);ttk.Button(pb,text='Открыть',command=lambda:self.app.popup_photo(self.current_path)).pack(side='left');row+=1
        act=ttk.Frame(form);act.grid(row=row,column=0,sticky='ew',pady=6);ttk.Button(act,text='Сохранить',command=self.save).pack(side='left');ttk.Button(act,text='Удалить карточку',command=self.delete).pack(side='left',padx=6);ttk.Button(act,text='Отмена/очистить',command=self.new).pack(side='left')
        form.columnconfigure(0,weight=1);form.rowconfigure(row-1,weight=1);self.status()

    def refresh(self):
        self.rows=self.db.search(self.q.get(),'Все поля',True)
        for i in self.tree.get_children():self.tree.delete(i)
        self.refresh_row_colors()
        for idx,r in enumerate(self.rows):self.tree.insert('', 'end',iid=str(r['id']),values=(idx+1,'│ '+safe_text(r['archive_no']),'│ '+safe_text(r['shot_date']),'│ '+safe_text(r['description'])[:100]),tags=('even' if idx % 2 == 0 else 'odd',))
        self.status_var.set(f'Записей: {len(self.rows)}')
    def select(self,_e=None):
        s=self.tree.selection();
        if not s:return
        r=self.db.get(int(s[0]));self.current_id=r['id'];self.current_path=r['file_path'] or ''
        for k,e in self.entries.items():e.delete(0,'end');e.insert(0,r[k] or '')
        self.desc.delete('1.0','end');self.desc.insert('1.0',r['description'] or '');self.show_preview()
    def new(self):
        self.current_id=None;self.current_path='';
        for e in self.entries.values():e.delete(0,'end')
        self.desc.delete('1.0','end');self.preview.configure(image='',text='Фото не выбрано');self.preview_ref=None
    def show_preview(self):
        self.preview_ref=None
        if not self.current_path or not os.path.exists(self.current_path) or Image is None:self.preview.configure(image='',text='Фото не выбрано');return
        try:
            img=Image.open(self.current_path);img.thumbnail((650,330));self.preview_ref=ImageTk.PhotoImage(img);self.preview.configure(image=self.preview_ref,text='')
        except Exception:self.preview.configure(image='',text='Ошибка изображения')
    def choose_photo(self):
        p=filedialog.askopenfilename(filetypes=[('Изображения','*.jpg *.jpeg *.png *.bmp *.tif *.tiff *.webp')]);
        if p:self.current_path=p;self.show_preview()
    def remove_photo(self):self.current_path='';self.show_preview()
    def collect(self):
        d={k:e.get().strip() for k,e in self.entries.items()};d['description']=self.desc.get('1.0','end').strip();d['file_path']=self.current_path;return d
    def save(self):
        d=self.collect();
        if not d['archive_no']:messagebox.showwarning(APP_NAME,'Укажите архивный номер.');return
        old=self.db.get(self.current_id) if self.current_id else None
        p=d['file_path']
        if p and os.path.exists(p):
            try:
                if old and old['file_path']==p:d['file_path']=p
                elif str(Path(p).resolve()).startswith(str(PHOTOS_DIR.resolve())):d['file_path']=p
                else:d['file_path']=copy_into_archive(p)
            except Exception as exc:messagebox.showerror(APP_NAME,str(exc));return
        if self.current_id:self.db.update(self.current_id,d)
        else:self.current_id=self.db.add(d)
        self.current_path=d['file_path'];self.refresh();self.status_var.set('Сохранено.')
    def delete(self):
        if not self.current_id:return
        if messagebox.askyesno(APP_NAME,'Удалить карточку из базы?'):
            self.db.delete(self.current_id);self.new();self.refresh()

    def _pick(self,row,aliases):
        norm={re.sub(r'[\s._№#:-]+',' ',safe_text(k).casefold()).strip():v for k,v in row.items()}
        for a in aliases:
            aa=re.sub(r'[\s._№#:-]+',' ',a.casefold()).strip()
            if aa in norm:return safe_text(norm[aa])
        return ''

    def _choose_import_mode(self):
        answer = simpledialog.askstring(
            APP_NAME,
            'Режим импорта:\n\n1 — добавить новые и обновить существующие по архивному №\n2 — добавить только новые\n3 — добавить все как новые карточки\n\nВведите 1, 2 или 3:',
            initialvalue='1', parent=self)
        if answer not in ('1','2','3'):
            return None
        return {'1':'upsert','2':'new_only','3':'append'}[answer]

    def _import_dict_rows(self, rows, base_dir, title='Импорт данных', mode='upsert'):
        aliases={
        'archive_no':['архивный номер','архивный №','арх №','арх. №','архивный номер фото','номер','archive no','archive_no','шифр'],
        'description':['описание','описание фотографии','содержание','аннотация','description','сюжет'],
        'shot_date':['дата съемки','дата съёмки','дата','дата фото','shot date','shot_date'],
        'location':['место съемки','место съёмки','место','location'],
        'author':['автор съемки','автор съёмки','автор','фотограф','author'],
        'source':['источник поступления','источник','source','поступление'],
        'file_path':['фото','фотография','путь к фото','файл','file path','file_path','имя файла']}
        prepared=[]; skipped_invalid=0
        total=len(rows)
        dlg=ProgressDialog(self, title)
        try:
            # Preparation pass. It intentionally does not refresh the main tables.
            for i,row in enumerate(rows,1):
                rec={k:self._pick(row,v) for k,v in aliases.items()}
                if not rec['archive_no']:
                    skipped_invalid += 1
                else:
                    fp=rec.get('file_path','')
                    if fp:
                        p=Path(fp)
                        if not p.is_absolute(): p=Path(base_dir)/p
                        # Copy linked images only when the referenced file really exists.
                        if p.exists() and p.suffix.lower() in IMAGE_EXTS:
                            try: rec['file_path']=copy_into_archive(p)
                            except Exception: rec['file_path']=''
                        else:
                            rec['file_path']=''
                    prepared.append(rec)
                if i == total or i % 100 == 0:
                    dlg.update_progress(i,total,f'Подготовка: {i} из {total}',f'Готово к записи: {len(prepared)} | Пропущено: {skipped_invalid}')

            added=updated=skipped=0
            if mode == 'upsert':
                def cb(i,t,a,u,sk):
                    dlg.update_progress(i,t,f'Запись в каталог: {i} из {t}',f'Добавлено: {a} | Обновлено: {u} | Пропущено: {sk+skipped_invalid}')
                added,updated,skipped=self.db.bulk_upsert(prepared,cb,batch_size=1000)
            else:
                cache=self.db.archive_index()
                t=len(prepared)
                self.db.conn.execute('BEGIN')
                try:
                    for i,rec in enumerate(prepared,1):
                        key=safe_text(rec['archive_no']).casefold()
                        if mode == 'new_only' and key in cache:
                            skipped += 1
                        else:
                            nid=self.db.add(rec,commit=False); added += 1
                            if mode == 'new_only':
                                cache.setdefault(key,[]).append(self.db.get(nid))
                        if i % 1000 == 0:
                            self.db.conn.commit(); self.db.conn.execute('BEGIN')
                        if i == t or i % 50 == 0:
                            dlg.update_progress(i,t,f'Запись в каталог: {i} из {t}',f'Добавлено: {added} | Пропущено: {skipped+skipped_invalid}')
                    self.db.conn.commit()
                except Exception:
                    self.db.conn.rollback(); raise
            skipped += skipped_invalid
            return added,updated,skipped
        finally:
            dlg.close()

    def import_table(self):
        path=filedialog.askopenfilename(title='Импорт Excel / CSV',filetypes=[('Excel/CSV','*.xlsx *.xlsm *.csv'),('Все файлы','*.*')])
        if not path:return
        mode=self._choose_import_mode()
        if not mode:return
        try:
            rows=[]
            if Path(path).suffix.lower() in ('.xlsx','.xlsm'):
                if load_workbook is None:raise RuntimeError('Не установлен openpyxl.')
                ws=load_workbook(path,read_only=True,data_only=True).active
                data=list(ws.iter_rows(values_only=True))
                if not data:raise ValueError('Таблица пустая.')
                headers=[safe_text(x) for x in data[0]]
                rows=[dict(zip(headers,v)) for v in data[1:] if any(x is not None for x in v)]
            else:
                try:
                    with open(path,'r',encoding='utf-8-sig',newline='') as f: rows=list(csv.DictReader(f))
                except UnicodeDecodeError:
                    with open(path,'r',encoding='cp1251',newline='') as f: rows=list(csv.DictReader(f))
            a,u,sk=self._import_dict_rows(rows,Path(path).parent,'Импорт Excel / CSV',mode)
            self.refresh();messagebox.showinfo(APP_NAME,f'Готово.\nДобавлено: {a}\nОбновлено: {u}\nПропущено: {sk}')
        except Exception as exc:messagebox.showerror(APP_NAME,f'Ошибка импорта:\n{exc}')

    def import_word(self):
        path=filedialog.askopenfilename(title='Импорт Word',filetypes=[('Word DOCX','*.docx'),('Все файлы','*.*')])
        if not path:return
        mode=self._choose_import_mode()
        if not mode:return
        try:
            if Document is None:raise RuntimeError('Не установлен python-docx.')
            doc=Document(path);rows=[]
            for table in doc.tables:
                if not table.rows:continue
                headers=[safe_text(c.text) for c in table.rows[0].cells]
                for tr in table.rows[1:]:
                    vals=[safe_text(c.text) for c in tr.cells]
                    if any(vals):rows.append(dict(zip(headers,vals)))
            if not rows:raise ValueError('В документе Word не найдено таблиц с данными.')
            a,u,sk=self._import_dict_rows(rows,Path(path).parent,'Импорт Word',mode)
            self.refresh();messagebox.showinfo(APP_NAME,f'Word импортирован.\nДобавлено: {a}\nОбновлено: {u}\nПропущено: {sk}')
        except Exception as exc:messagebox.showerror(APP_NAME,f'Ошибка импорта Word:\n{exc}')

    def import_clipboard(self):
        try:
            text=self.clipboard_get()
        except Exception:
            messagebox.showinfo(APP_NAME,'В буфере обмена нет текстовой таблицы.');return
        lines=[ln for ln in text.splitlines() if ln.strip()]
        if len(lines)<2:
            messagebox.showinfo(APP_NAME,'Скопируйте из Word или Excel таблицу вместе со строкой заголовков.');return
        rows_raw=[ln.split('\t') for ln in lines]
        headers=[safe_text(x) for x in rows_raw[0]]
        if len(headers)<2:
            messagebox.showinfo(APP_NAME,'Не удалось распознать столбцы. Скопируйте таблицу из Word/Excel целиком.');return
        rows=[dict(zip(headers,r+['']*max(0,len(headers)-len(r)))) for r in rows_raw[1:]]
        mode=self._choose_import_mode()
        if not mode:return
        try:
            a,u,sk=self._import_dict_rows(rows,DATA_DIR,'Вставка из буфера',mode)
            self.refresh();messagebox.showinfo(APP_NAME,f'Вставка завершена.\nДобавлено: {a}\nОбновлено: {u}\nПропущено: {sk}')
        except Exception as exc:messagebox.showerror(APP_NAME,f'Ошибка вставки:\n{exc}')

    def import_access(self):
        path=filedialog.askopenfilename(title='Импорт Microsoft Access',filetypes=[('Access','*.accdb *.mdb')])
        if not path:return
        mode=self._choose_import_mode()
        if not mode:return
        try:
            import pyodbc
            drivers=[d for d in pyodbc.drivers() if 'Access Driver' in d]
            if not drivers:raise RuntimeError('Не установлен Microsoft Access Database Engine (ODBC).')
            conn=pyodbc.connect(r'DRIVER={'+drivers[-1]+r'};DBQ='+path+';')
            tables=[r.table_name for r in conn.cursor().tables(tableType='TABLE') if not r.table_name.startswith('MSys')]
            if not tables:raise RuntimeError('Таблицы не найдены.')
            table=tables[0]
            if len(tables)>1:
                table=simpledialog.askstring(APP_NAME,'Таблицы: '+', '.join(tables[:15])+'\nВведите имя таблицы:',initialvalue=tables[0]) or ''
                if table not in tables:return
            cur=conn.cursor()
            cur.execute(f'SELECT COUNT(*) FROM [{table}]'); total=int(cur.fetchone()[0] or 0)
            cur.execute(f'SELECT * FROM [{table}]'); cols=[d[0] for d in cur.description]
            rows=[]; dlg=ProgressDialog(self,'Чтение Microsoft Access')
            try:
                done=0
                while True:
                    batch=cur.fetchmany(1000)
                    if not batch:break
                    rows.extend(dict(zip(cols,r)) for r in batch)
                    done += len(batch)
                    dlg.update_progress(done,max(total,done),f'Прочитано: {done} из {total or "?"}',f'Таблица: {table}')
            finally:
                dlg.close(); conn.close()
            a,u,sk=self._import_dict_rows(rows,Path(path).parent,'Быстрый импорт Access',mode)
            self.refresh();messagebox.showinfo(APP_NAME,f'Таблица: {table}\nДобавлено: {a}\nОбновлено: {u}\nПропущено: {sk}')
        except Exception as exc:messagebox.showerror(APP_NAME,f'Ошибка Access:\n{exc}')

    def _match_archive_for_filename(self,stem,index):
        key=normalize_key(stem)
        if key in index:return index[key]
        # Common forms: ARCHIVENO_1, ARCHIVENO(2), ARCHIVENO-copy
        candidates=[]
        for k,rows in index.items():
            if not k:continue
            if key.startswith(k) or k in key:candidates.append((len(k),rows))
        if not candidates:return None
        candidates.sort(key=lambda x:x[0],reverse=True);return candidates[0][1]

    def import_photos_by_number(self):
        folder=filedialog.askdirectory(title='Выберите папку с фотографиями')
        if not folder:return
        recursive=messagebox.askyesno(APP_NAME,'Искать фотографии также во вложенных папках?')
        files=[]
        it=Path(folder).rglob('*') if recursive else Path(folder).glob('*')
        for p in it:
            if p.is_file() and p.suffix.lower() in IMAGE_EXTS:files.append(p)
        if not files:messagebox.showinfo(APP_NAME,'Фотографии не найдены.');return
        index={}
        for r in self.db.all():index.setdefault(normalize_key(r['archive_no']),[]).append(r)
        attached=cloned=unmatched=errors=0
        dlg=ProgressDialog(self,'Импорт фотографий по архивному №')
        try:
            total=len(files)
            for n,p in enumerate(files,1):
                matches=self._match_archive_for_filename(p.stem,index)
                if not matches:
                    unmatched+=1
                else:
                    try:
                        target=next((r for r in matches if not r['file_path']),None)
                        dest=copy_into_archive(p)
                        if target:
                            rec={k:target[k] for k in CatalogDB.FIELDS};rec['file_path']=dest
                            self.db.update(target['id'],rec);attached+=1
                            # refresh local index row
                            updated=self.db.get(target['id'])
                            for j,rr in enumerate(matches):
                                if rr['id']==target['id']:matches[j]=updated;break
                        else:
                            src=matches[0];rec={k:src[k] for k in CatalogDB.FIELDS};rec['file_path']=dest
                            nid=self.db.add(rec);cloned+=1
                            matches.append(self.db.get(nid))
                    except Exception:
                        errors+=1
                if n==total or n%10==0:
                    dlg.update_progress(n,total,f'Файл: {p.name}',f'Привязано: {attached} | Доп. карточек: {cloned} | Не найдено: {unmatched} | Ошибок: {errors}')
        finally:
            dlg.close()
        self.refresh();messagebox.showinfo(APP_NAME,f'Импорт фотографий завершён.\nПривязано к пустым карточкам: {attached}\nДополнительных карточек: {cloned}\nНе найден архивный номер: {unmatched}\nОшибок: {errors}')

    def index_faces(self):
        if not self.app.face.available:messagebox.showerror(APP_NAME,'Поиск по лицу недоступен:\n'+self.app.face.error);return
        rows=self.db.all();photos=faces=noface=errors=0
        dlg=ProgressDialog(self,'Индексация лиц')
        try:
            total=len(rows)
            for i,r in enumerate(rows,1):
                try:
                    if r['file_path'] and os.path.exists(r['file_path']):
                        embs=self.app.face.embeddings(r['file_path']);self.db.replace_faces(r['id'],embs);photos+=1;faces+=len(embs)
                    else:errors+=1
                except ValueError:self.db.replace_faces(r['id'],[]);noface+=1
                except Exception:errors+=1
                if i==total or i%5==0:
                    dlg.update_progress(i,total,f'Обработано: {i} из {total}',f'Фото: {photos} | Лиц: {faces} | Без лиц: {noface} | Ошибок: {errors}')
        finally:
            dlg.close()
        messagebox.showinfo(APP_NAME,f'Индексация завершена.\nФотографий обработано: {photos}\nЛиц проиндексировано: {faces}\nБез лиц: {noface}\nОшибок: {errors}')
        self.status_var.set('Готово')

    def backup(self):
        target=filedialog.asksaveasfilename(defaultextension='.zip',initialfile=f"PhotoArchive_backup_{datetime.now():%Y%m%d_%H%M}.zip",filetypes=[('ZIP','*.zip')]);
        if not target:return
        try:
            self.db.conn.commit()
            with zipfile.ZipFile(target,'w',zipfile.ZIP_DEFLATED) as z:
                for p in DATA_DIR.rglob('*'):
                    if p.is_file():z.write(p,p.relative_to(DATA_DIR))
            messagebox.showinfo(APP_NAME,'Резервная копия создана.')
        except Exception as exc:messagebox.showerror(APP_NAME,str(exc))


class CatalogScreen(BaseScreen):
    def __init__(self,parent,app):
        super().__init__(parent,app)
        self.nav('catalog')
        self.thumb_refs={}
        self.thumb_size=tk.IntVar(value=72)
        self.page_size=tk.IntVar(value=200)
        self.page=0
        self.all_rows=[]
        self.checked=set()

        top1=ttk.Frame(self,padding=(10,0,10,4));top1.pack(fill='x')
        self.q=tk.StringVar()
        ttk.Label(top1,text='Поиск:').pack(side='left')
        e=ttk.Entry(top1,textvariable=self.q,width=34);e.pack(side='left',padx=5);e.bind('<Return>',lambda x:self.new_search())
        ttk.Button(top1,text='Найти',command=self.new_search).pack(side='left')
        ttk.Button(top1,text='Импорт Word',command=self.catalog_import_word).pack(side='left',padx=(12,4))
        ttk.Button(top1,text='Вставить из Word/Excel',command=self.catalog_import_clipboard).pack(side='left',padx=4)
        ttk.Button(top1,text='Импорт фото по архивным №',command=self.catalog_import_photos).pack(side='left',padx=4)
        ttk.Label(top1,text='Миниатюры:').pack(side='right')
        ttk.Scale(top1,from_=48,to=110,variable=self.thumb_size,command=self.resize_thumbs).pack(side='right',padx=6)

        top2=ttk.Frame(self,padding=(10,0,10,6));top2.pack(fill='x')
        ttk.Button(top2,text='Выбрать страницу',command=self.select_all).pack(side='left')
        ttk.Button(top2,text='Снять выбор',command=self.clear_selection).pack(side='left',padx=4)
        ttk.Button(top2,text='Удалить выбранные',command=self.delete_selected).pack(side='left',padx=4)
        ttk.Separator(top2,orient='vertical').pack(side='left',fill='y',padx=10)
        ttk.Button(top2,text='◀ Назад',command=self.prev_page).pack(side='left')
        ttk.Button(top2,text='Далее ▶',command=self.next_page).pack(side='left',padx=4)
        ttk.Label(top2,text='На странице:').pack(side='left',padx=(12,3))
        cb=ttk.Combobox(top2,textvariable=self.page_size,state='readonly',width=6,values=(100,200,500))
        cb.pack(side='left');cb.bind('<<ComboboxSelected>>',lambda e:self.change_page_size())
        self.page_var=tk.StringVar(value='')
        ttk.Label(top2,textvariable=self.page_var,font=('Arial',10,'bold')).pack(side='left',padx=14)

        fr=ttk.Frame(self,padding=(10,0,10,8));fr.pack(fill='both',expand=True)
        cols=('mark','seq','no','desc','date','loc','author','source')
        self.tree=ttk.Treeview(fr,columns=cols,show='tree headings',selectmode='extended',style='Catalog.Treeview')
        self.tree.heading('#0',text='Фото');self.tree.column('#0',width=105,minwidth=90,stretch=False)
        self.tree.heading('mark',text='✓'); self.tree.column('mark',width=44,minwidth=44,stretch=False,anchor='center')
        self.tree.heading('seq',text='№'); self.tree.column('seq',width=65,minwidth=55,stretch=False,anchor='center')
        specs=[('no','Архивный №',135),('desc','Описание',430),('date','Дата',105),('loc','Место',160),('author','Автор',160),('source','Источник',220)]
        for c,t,w in specs:self.tree.heading(c,text=t);self.tree.column(c,width=w,anchor='w')
        ys=ttk.Scrollbar(fr,orient='vertical',command=self.tree.yview);xs=ttk.Scrollbar(fr,orient='horizontal',command=self.tree.xview)
        self.tree.configure(yscrollcommand=ys.set,xscrollcommand=xs.set)
        self.tree.grid(row=0,column=0,sticky='nsew');ys.grid(row=0,column=1,sticky='ns');xs.grid(row=1,column=0,sticky='ew')
        fr.rowconfigure(0,weight=1);fr.columnconfigure(0,weight=1)
        self.tree.bind('<Double-1>',self.open_clicked)
        self.tree.bind('<ButtonRelease-1>',self.click_open)
        self.tree.bind('<Motion>',self.hover)
        self.tree.bind('<Leave>',self.hide_hover)
        self.tree.bind('<Delete>',lambda e:self.delete_selected())
        self.tree.bind('<Control-a>',self.ctrl_a)
        self.status()

    def catalog_import_word(self):
        self.app.screens['editor'].import_word(); self.new_search()

    def catalog_import_clipboard(self):
        self.app.screens['editor'].import_clipboard(); self.new_search()

    def catalog_import_photos(self):
        self.app.screens['editor'].import_photos_by_number(); self.new_search()

    def new_search(self):
        self.page=0; self.refresh()

    def change_page_size(self):
        self.page=0; self.refresh()

    def refresh(self):
        self.all_rows=self.db.search(self.q.get(),'Все поля',True)
        total=len(self.all_rows)
        ps=max(1,int(self.page_size.get()))
        pages=max(1, math.ceil(total/ps))
        self.page=min(max(0,self.page),pages-1)
        start=self.page*ps; end=min(total,start+ps)
        self.rows=self.all_rows[start:end]
        for i in self.tree.get_children():self.tree.delete(i)
        self.thumb_refs={}
        size=max(48,int(self.thumb_size.get()));self.app.style.configure('Catalog.Treeview',rowheight=max(58,size+8))
        self.refresh_row_colors()
        valid_ids={str(r['id']) for r in self.all_rows}; self.checked={i for i in self.checked if i in valid_ids}
        for local_idx,r in enumerate(self.rows):
            global_idx=start+local_idx+1
            tp=make_thumb(r['file_path'],r['id'],(size,int(size*.75)));img=''
            if tp and ImageTk:
                try:img=ImageTk.PhotoImage(Image.open(tp));self.thumb_refs[str(r['id'])]=img
                except Exception:img=''
            iid=str(r['id']); mark='☑' if iid in self.checked else '☐'
            self.tree.insert('', 'end',iid=iid,text='',image=img,
                             values=(mark,global_idx,'│ '+safe_text(r['archive_no']),'│ '+safe_text(r['description']).replace('\n',' ')[:180],'│ '+safe_text(r['shot_date']),'│ '+safe_text(r['location']),'│ '+safe_text(r['author']),'│ '+safe_text(r['source'])),
                             tags=('even' if local_idx % 2 == 0 else 'odd',))
        shown=f'{start+1}–{end}' if total else '0'
        self.page_var.set(f'Страница {self.page+1} из {pages}')
        self.status_var.set(f'Записей в базе/поиске: {total} | Показано: {shown}')

    def prev_page(self):
        if self.page>0:self.page-=1;self.refresh()

    def next_page(self):
        ps=max(1,int(self.page_size.get()))
        if (self.page+1)*ps < len(self.all_rows):self.page+=1;self.refresh()

    def ctrl_a(self,_e=None):
        self.select_all();return 'break'

    def select_all(self):
        ids=list(self.tree.get_children())
        self.tree.selection_set(ids);self.checked.update(ids)
        for iid in ids:
            vals=list(self.tree.item(iid,'values'))
            if vals:vals[0]='☑';self.tree.item(iid,values=vals)
        self.status_var.set(f'Выбрано записей: {len(self.checked)}')

    def clear_selection(self):
        self.tree.selection_remove(self.tree.selection());self.checked.clear()
        for iid in self.tree.get_children():
            vals=list(self.tree.item(iid,'values'))
            if vals:vals[0]='☐';self.tree.item(iid,values=vals)
        self.status_var.set(f'Выбор снят. Всего найдено: {len(self.all_rows)}')

    def delete_selected(self):
        ids=set(self.tree.selection()) | set(self.checked)
        if not ids:
            messagebox.showinfo(APP_NAME,'Сначала выберите одну или несколько записей.'); return
        if not messagebox.askyesno(APP_NAME,f'Удалить выбранные записи из базы?\nКоличество: {len(ids)}'):
            return
        choice=messagebox.askyesnocancel(APP_NAME,'Удалить также файлы фотографий с диска?\n\nДа — удалить записи и файлы\nНет — удалить только записи\nОтмена — ничего не удалять')
        if choice is None:return
        delete_files=bool(choice);removed=0;file_errors=0
        for iid in list(ids):
            try:
                row=self.db.delete(int(iid)); removed+=1
                if delete_files and row and row['file_path']:
                    try:
                        p=Path(row['file_path'])
                        if p.exists() and str(p.resolve()).startswith(str(PHOTOS_DIR.resolve())):p.unlink()
                    except Exception:file_errors+=1
            except Exception:pass
        self.checked.clear();self.refresh()
        msg=f'Удалено записей: {removed}'
        if file_errors:msg+=f'\nНе удалось удалить файлов: {file_errors}'
        messagebox.showinfo(APP_NAME,msg)

    def resize_thumbs(self,_=None):
        if hasattr(self,'_resize_job'):
            try:self.after_cancel(self._resize_job)
            except Exception:pass
        self._resize_job=self.after(250,self.refresh)

    def _row_path_at(self,y):
        iid=self.tree.identify_row(y)
        if not iid:return None,None
        r=self.db.get(int(iid));return r,r['file_path'] if r else None

    def click_open(self,e):
        region=self.tree.identify_region(e.x,e.y);col=self.tree.identify_column(e.x);iid=self.tree.identify_row(e.y)
        if iid and col=='#1':
            if iid in self.checked:self.checked.remove(iid)
            else:self.checked.add(iid)
            vals=list(self.tree.item(iid,'values'))
            if vals:vals[0]='☑' if iid in self.checked else '☐';self.tree.item(iid,values=vals)
            self.status_var.set(f'Отмечено галочками: {len(self.checked)}')
            return
        if region=='tree' or col=='#0':
            r,p=self._row_path_at(e.y)
            if p:self.app.popup_photo(p,r['archive_no'])

    def open_clicked(self,e):
        r,p=self._row_path_at(e.y)
        if p:self.app.popup_photo(p,r['archive_no'])

    def hover(self,e):
        # Hover preview only when the mouse is over the thumbnail column.
        if self.tree.identify_column(e.x)!='#0':return self.hide_hover()
        r,p=self._row_path_at(e.y)
        if not p or not os.path.exists(p) or Image is None:return self.hide_hover()
        if self.hover_after:
            try:self.after_cancel(self.hover_after)
            except Exception:pass
        self.hover_after=self.after(350,lambda:self.show_hover(p,e.x_root+18,e.y_root+18))

    def show_hover(self,p,x,y):
        self.hide_hover()
        try:
            img=Image.open(p);target_h=190;ratio=target_h/max(1,img.height);img=img.resize((max(1,int(img.width*ratio)),target_h))
            top=tk.Toplevel(self);top.overrideredirect(True);top.geometry(f'+{x}+{y}')
            ref=ImageTk.PhotoImage(img);lab=ttk.Label(top,image=ref,relief='solid');lab.image=ref;lab.pack();self.app.hover_popup=top
        except Exception:pass

    def hide_hover(self,_=None,cancel_only=False):
        if self.hover_after:
            try:self.after_cancel(self.hover_after)
            except Exception:pass
            self.hover_after=None
        if not cancel_only and self.app.hover_popup:
            try:self.app.hover_popup.destroy()
            except Exception:pass
            self.app.hover_popup=None


if __name__=='__main__':
    PhotoCatalogApp().mainloop()
